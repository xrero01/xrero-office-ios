"""make_samples.py - realistic sample files for the App Store screenshots (Arabic + English).
Writes store/samples/{report_ar.docx, report_en.docx, sales_ar.xlsx, sales_en.xlsx, deck_ar.pptx, deck_en.pptx}."""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt as PPt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
os.makedirs(OUT, exist_ok=True)
TEAL, DEEP, GOLD, GREY = (0x1F, 0x7A, 0x8C), (0x15, 0x4F, 0x5E), (0xC8, 0x8A, 0x12), (0x55, 0x5F, 0x66)

# ------------------------------------------------------------------ Word
def rtl_paragraph(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:bidi"); pPr.append(b)

def run(p, text, size=11, bold=False, color=None, font="Cairo", rtl=False):
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
    rPr = r._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), font)
    if rtl:
        e = OxmlElement("w:rtl"); rPr.append(e)
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rPr.append(szcs)
        if bold: rPr.append(OxmlElement("w:bCs"))
    if color: r.font.color.rgb = RGBColor(*color)
    return r

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)

def report(lang):
    ar = lang == "ar"
    font = "Cairo" if ar else "Calibri"
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = Cm(2)
    T = {
        "title": ("تقرير الأداء الربعي", "Quarterly Performance Report"),
        "sub": ("الربع الثالث 2026 · فروع الإمارات", "Q3 2026 · UAE branches"),
        "intro": ("حققت الشركة نمواً قوياً هذا الربع، مدفوعاً بارتفاع المبيعات في فروع دبي وأبوظبي وتحسّن كفاءة التشغيل.",
                  "The company grew strongly this quarter, driven by higher sales in Dubai and Abu Dhabi and leaner operations."),
        "h1": ("أبرز النتائج", "Highlights"),
        "b": (["ارتفعت الإيرادات بنسبة 18% مقارنة بالربع السابق", "انخفضت التكاليف التشغيلية بنسبة 6%", "افتتاح فرع جديد في الشارقة"],
              ["Revenue up 18% on the previous quarter", "Operating costs down 6%", "New branch opened in Sharjah"]),
        "h2": ("الإيرادات حسب الفرع", "Revenue by branch"),
        "cols": (["الفرع", "الإيرادات (درهم)", "النمو"], ["Branch", "Revenue (AED)", "Growth"]),
        "rows": ([["دبي", "1,240,000", "+21%"], ["أبوظبي", "980,000", "+17%"], ["الشارقة", "410,000", "+9%"]],
                 [["Dubai", "1,240,000", "+21%"], ["Abu Dhabi", "980,000", "+17%"], ["Sharjah", "410,000", "+9%"]]),
        "end": ("نشكر جميع فرق العمل على جهودهم المتميزة، ونتطلع إلى ربع رابع أقوى.",
                "Thank you to every team for an outstanding quarter - we look forward to an even stronger Q4."),
    }
    k = 0 if ar else 1
    align = WD_ALIGN_PARAGRAPH.RIGHT if ar else WD_ALIGN_PARAGRAPH.LEFT
    def para(text, size=11, bold=False, color=None, space=6):
        p = d.add_paragraph(); p.alignment = align
        if ar: rtl_paragraph(p)
        run(p, text, size, bold, color, font, ar); p.paragraph_format.space_after = Pt(space)
        return p
    para(T["title"][k], 24, True, DEEP, 2)
    para(T["sub"][k], 12, False, TEAL, 14)
    para(T["intro"][k], 11.5, False, None, 12)
    para(T["h1"][k], 15, True, TEAL, 6)
    for b in T["b"][k]:
        p = para("•  " + b, 11.5, False, None, 3)
    para("", 6)
    para(T["h2"][k], 15, True, TEAL, 6)
    t = d.add_table(rows=1 + len(T["rows"][k]), cols=3)
    t.style = "Table Grid"
    if ar:
        tblPr = t._tbl.tblPr; bv = OxmlElement("w:bidiVisual"); tblPr.append(bv)
    for c, h in enumerate(T["cols"][k]):
        cell = t.rows[0].cells[c]; cell.text = ""; p = cell.paragraphs[0]; p.alignment = align
        if ar: rtl_paragraph(p)
        run(p, h, 11, True, (255, 255, 255), font, ar); shade(cell, "1F7A8C")
    for r_i, row in enumerate(T["rows"][k], start=1):
        for c, v in enumerate(row):
            cell = t.rows[r_i].cells[c]; cell.text = ""; p = cell.paragraphs[0]; p.alignment = align
            if ar: rtl_paragraph(p)
            run(p, v, 11, c == 0, (0x1B, 0x8A, 0x4B) if c == 2 else None, font, ar)
            if r_i % 2 == 0: shade(cell, "EEF6F8")
    para("", 6)
    para(T["end"][k], 11.5, False, GREY, 6)
    d.save(os.path.join(OUT, "report_%s.docx" % lang))

# ------------------------------------------------------------------ Excel
def sales(lang):
    ar = lang == "ar"
    wb = Workbook(); ws = wb.active
    ws.title = "المبيعات" if ar else "Sales"
    ws.sheet_view.rightToLeft = ar
    font = "Cairo" if ar else "Calibri"
    months = (["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو"] if ar else ["Jan", "Feb", "Mar", "Apr", "May", "Jun"])
    branches = (["دبي", "أبوظبي", "الشارقة"] if ar else ["Dubai", "Abu Dhabi", "Sharjah"])
    data = [[312, 298, 341, 365, 388, 402], [254, 262, 271, 290, 301, 318], [98, 104, 112, 121, 133, 140]]
    ws["A1"] = "مبيعات النصف الأول 2026 (ألف درهم)" if ar else "H1 2026 sales (AED thousands)"
    ws["A1"].font = Font(name=font, size=15, bold=True, color="154F5E")
    head = ["الفرع" if ar else "Branch"] + months + ["الإجمالي" if ar else "Total"]
    thin = Side(style="thin", color="C9D6DA")
    for c, h in enumerate(head, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(name=font, bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F7A8C")
        cell.alignment = Alignment(horizontal="center"); cell.border = Border(bottom=thin)
    for r, (b, vals) in enumerate(zip(branches, data), start=4):
        ws.cell(row=r, column=1, value=b).font = Font(name=font, bold=True)
        for c, v in enumerate(vals, start=2):
            cell = ws.cell(row=r, column=c, value=v); cell.font = Font(name=font); cell.alignment = Alignment(horizontal="center")
        tot = ws.cell(row=r, column=8, value="=SUM(B%d:G%d)" % (r, r)); tot.font = Font(name=font, bold=True, color="154F5E")
        tot.alignment = Alignment(horizontal="center")
        if r % 2 == 1:
            for c in range(1, 9): ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="EEF6F8")
    ws.cell(row=7, column=1, value="المجموع" if ar else "Total").font = Font(name=font, bold=True)
    for c in range(2, 9):
        col = chr(64 + c)
        cell = ws.cell(row=7, column=c, value="=SUM(%s4:%s6)" % (col, col))
        cell.font = Font(name=font, bold=True); cell.alignment = Alignment(horizontal="center"); cell.border = Border(top=thin)
    ws.column_dimensions["A"].width = 14
    for c in "BCDEFGH": ws.column_dimensions[c].width = 10
    ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
    ch.title = "المبيعات الشهرية" if ar else "Monthly sales"
    ch.add_data(Reference(ws, min_col=1, max_col=7, min_row=4, max_row=6), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=7, min_row=3))
    ch.height, ch.width = 7.5, 15
    ws.add_chart(ch, "A9")
    wb.save(os.path.join(OUT, "sales_%s.xlsx" % lang))

# ------------------------------------------------------------------ PowerPoint
def deck(lang):
    ar = lang == "ar"
    font = "Cairo" if ar else "Calibri"
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = PRGB(*DEEP)
    band = s.shapes.add_shape(1, 0, Emu(4900000), prs.slide_width, Emu(1958000)); band.fill.solid(); band.fill.fore_color.rgb = PRGB(*TEAL); band.line.fill.background()
    def text(x, y, w, h, t, size, bold=False, color=(255, 255, 255)):
        tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT if ar else PP_ALIGN.LEFT
        r = p.add_run(); r.text = t; r.font.size = PPt(size); r.font.bold = bold; r.font.name = font; r.font.color.rgb = PRGB(*color)
        if ar:
            pPr = p._p.get_or_add_pPr(); pPr.set("rtl", "1")
    text(900000, 1500000, 10400000, 1400000, "خطة النمو 2027" if ar else "Growth Plan 2027", 54, True)
    text(900000, 2900000, 10400000, 800000, "ثلاث أولويات لعام أقوى" if ar else "Three priorities for a stronger year", 26, False, (0xD8, 0xEE, 0xF2))
    text(900000, 5300000, 10400000, 700000, "فريق الإدارة · أكتوبر 2026" if ar else "Leadership team · October 2026", 18, False, (255, 255, 255))
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s = s2
    text(900000, 500000, 10400000, 900000, "الأولويات" if ar else "Priorities", 36, True, DEEP)
    items = (["توسيع الفروع في الإمارات الشمالية", "إطلاق المتجر الإلكتروني", "رفع رضا العملاء إلى 95%"] if ar
             else ["Expand into the Northern Emirates", "Launch the online store", "Lift customer satisfaction to 95%"])
    for i, it in enumerate(items):
        text(900000, 1700000 + i * 1100000, 10400000, 900000, "●  " + it, 26, False, (0x33, 0x3D, 0x42))
    prs.save(os.path.join(OUT, "deck_%s.pptx" % lang))

for lang in ("ar", "en"):
    report(lang); sales(lang); deck(lang)
print(sorted(os.listdir(OUT)))
